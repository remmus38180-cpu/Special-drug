"""驗收案例 T01–T40（規格書第14節）。測試名稱即案例編號。"""
import json
import os
import subprocess
import sys
import time
from pathlib import Path

import pytest

from conftest import BASE_ROWS, HEAD, MAPPING, write_master
from specdrug import cases as C
from specdrug import master_import as M
from specdrug import migration as MG
from specdrug import pricing as P
from specdrug import recognition as R
from specdrug import rules
from specdrug.attachments import add_attachment
from specdrug.core import BizError, Ctx, PermissionDenied, ROLE_VIEWER
from specdrug.query import drug_detail
from specdrug.store import LockBusy, LockLost, PublishConflict, SharedStore, StoreError, verify_backup
from specdrug.workspace import Workspace

ROOT = Path(__file__).resolve().parents[1]
UPLOAD = Path("/root/.claude/uploads/c68f97b3-18c4-5255-aae1-4db861d58cf0/29a98e42-___1.xlsx")


# ---------------- helpers ----------------

def recognized(ctx, imp, rule, *, name="認定", category="特殊藥品", eff="2020-01-01", confirm=True,
               decide_pending=None):
    rec = R.create_recognition(ctx, name=name)
    rev = R.create_revision(ctx, rec, category=category, rule=rule, effective_from=eff, completeness="完整")
    sid = R.save_snapshot(ctx, rev, imp)
    for m in ctx.con.execute("SELECT drug_code FROM scope_member WHERE snapshot_id=? AND eval_result='待確認'", (sid,)):
        R.set_member_decision(ctx, sid, m[0], decide_pending or "排除", "測試人工確認")
    R.confirm_snapshot(ctx, sid)
    if confirm:
        R.confirm_revision(ctx, rev)
    return rec, rev, sid


def signed_contract(ctx, codes, expiry="2099-01-01"):
    return P.create_contract(ctx, contract_no="C-001", firm="甲藥廠", status="完成簽署", signed_date="2026-01-01",
                             effective_from="2026-01-01", expiry_date=expiry, drugs=codes)


def ready_action(ctx, rev, code, **kw):
    args = dict(revision_id=rev, drug_code=code, contract_requirement="當時無須簽約", requirement_basis="歷史規定",
                price_after="150", effective_date="2026-11-01", announce_doc_no="健保審字第1150000001號",
                announce_date="2026-10-20", status="已核定")
    args.update(kw)
    return P.create_price_action(ctx, **args)


def new_store(tmp_path, account="admin"):
    s = SharedStore(tmp_path / "share")
    s.init(account)
    return s


# ---------------- 主檔 ----------------

def test_T01_leading_zero_and_letter_codes(ctx, tmp_path):
    rows = BASE_ROWS + [["1.23E+10", "SCI", "", "", "", "", "", "", "", "", "", ""]]
    iid = M.stage_import(ctx, write_master(tmp_path / "a.csv", rows), MAPPING, base_date="2026-10-01")
    codes = {r[0] for r in ctx.con.execute("SELECT drug_code FROM master_drug WHERE import_id=?", (iid,))}
    assert "0A00005100" in codes and "B000004100" in codes
    chk = M.check_of(ctx.con, iid)
    assert chk["code_error_count"] == 1 and "科學記號" in chk["code_errors"][0]["error"]
    # xlsx數值型代碼：列為錯誤，不自行補碼
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["CODE", "藥品名稱"])
    ws.append([12345, "NUM"])
    ws.append(["0012345", "TXT"])
    wb.save(tmp_path / "n.xlsx")
    iid2 = M.stage_import(ctx, tmp_path / "n.xlsx", {"drug_code": "CODE", "drug_name": "藥品名稱"})
    chk2 = M.check_of(ctx.con, iid2)
    assert chk2["code_error_count"] == 1 and "前導零" in chk2["code_errors"][0]["error"]
    assert ctx.con.execute("SELECT drug_code FROM master_drug WHERE import_id=?", (iid2,)).fetchone()[0] == "0012345"
    with pytest.raises(BizError):
        M.activate(ctx, iid2)


def test_T02_conflicting_duplicate_blocks_activation(ctx, tmp_path):
    rows = BASE_ROWS + [["A000001100", "DRUG A1", "甲藥廠", "G1AAA0000001", "J06BA01", "20", "注射劑", "N", "5",
                         "ML", "999", "0"]]
    iid = M.stage_import(ctx, write_master(tmp_path / "b.csv", rows), MAPPING, base_date="2026-10-01")
    with pytest.raises(BizError) as e:
        M.activate(ctx, iid)
    assert "A000001100" in str(e.value) and "列 2,9" in str(e.value) and "price" in str(e.value)
    assert M.active_import_id(ctx.con) is None
    # 完全相同重複：確認後才可合併啟用
    rows2 = BASE_ROWS + [BASE_ROWS[0]]
    iid2 = M.stage_import(ctx, write_master(tmp_path / "c.csv", rows2), MAPPING, base_date="2026-10-01")
    with pytest.raises(BizError):
        M.activate(ctx, iid2)
    M.activate(ctx, iid2, confirm_merge_duplicates=True)
    assert M.active_import_id(ctx.con) == iid2


def test_T03_two_groups_union(ctx, master):
    rule = {"field": "group_code", "op": "in", "values": ["G1AAA0000001", "G1CCC0000001"]}
    pv = R.preview(ctx.con, rule, master)
    codes = [m["drug_code"] for m in pv["members"]]
    assert sorted(codes) == ["A000001100", "A000002100", "A000006100", "B000004100"]
    assert len(codes) == len(set(codes))


def test_T04_atc_multi_dose_and(ctx, master):
    rule = {"op": "and", "children": [
        {"field": "atc_code", "op": "in", "values": ["L04AA03", "L04AA04", "C01AA05"]},
        {"field": "multi_mark", "op": "eq", "value": "N"},
        {"field": "dose_code", "op": "eq", "value": "11"}]}
    pv = R.preview(ctx.con, rule, master)
    res = {m["drug_code"]: m["result"] for m in pv["members"]}
    assert res == {"B000004100": "符合", "A000007100": "待確認"}  # D1為複方→不符合；F1缺單複方→待確認


def test_T05_prefix_exclude_or(ctx, master):
    rule = {"op": "or", "children": [
        {"op": "and", "children": [{"field": "atc_code", "op": "starts_with", "value": "J06B"},
                                   {"field": "atc_code", "op": "not_in", "values": ["J06BB16"]}]},
        {"field": "atc_code", "op": "in", "values": ["L04AA03", "L04AA04"]}]}
    pv = R.preview(ctx.con, rule, master)
    codes = sorted(m["drug_code"] for m in pv["members"] if m["result"] == "符合")
    assert codes == ["0A00005100", "A000001100", "A000002100", "A000006100", "B000004100"]
    assert "A000003100" not in codes
    # 原文解析得到相同邏輯
    ast, issues = rules.parse_rule_text("(ATC前4碼=J06B & ATC7碼<>J06BB16) or ATC7碼=L04AA03  or ATC7碼=L04AA04")
    assert not issues
    assert sorted(m["drug_code"] for m in R.preview(ctx.con, ast, master)["members"]) == codes


def test_T06_group_plus_spec(ctx, master):
    rule = {"op": "and", "children": [{"field": "group_code", "op": "eq", "value": "G1AAA0000001"},
                                      {"field": "spec", "op": "eq", "value": "5", "unit": "mL"}]}
    pv = R.preview(ctx.con, rule, master)
    res = {m["drug_code"]: m["result"] for m in pv["members"]}
    assert res == {"A000001100": "符合", "A000006100": "待確認"}  # 10ML排除；XYZ未知單位待確認
    g = pv["groups"][0]
    assert g["partial"] is True and any("部分涵蓋分組" in w for w in pv["warnings"])


def test_T07_cross_item_condition_requires_manual(ctx, master):
    ast, issues = rules.parse_rule_text(
        "分組代碼= B01WAR212112，且同廠牌尚有不同規格品項，﹛分組代碼前11碼=B01WAR21211 & 分組代碼<>B01WAR21211A)=0﹜")
    assert ast["op"] == "manual" and any("=0" in i for i in issues) and any("同廠牌" in i for i in issues)
    rule = {"op": "and", "children": [{"field": "group_code", "op": "eq", "value": "G1AAA0000001"},
                                      {"op": "manual", "text": "同廠牌尚有不同規格品項"}]}
    rec = R.create_recognition(ctx, name="Warfarin")
    rev = R.create_revision(ctx, rec, category="不可替代特殊藥品", rule=rule, effective_from="2020-01-01")
    sid = R.save_snapshot(ctx, rev, master)
    assert {r[0] for r in ctx.con.execute("SELECT eval_result FROM scope_member WHERE snapshot_id=?", (sid,))} == {"待確認"}
    with pytest.raises(BizError, match="待確認品項未人工確認"):
        R.confirm_snapshot(ctx, sid)
    with pytest.raises(BizError, match="尚未確認範圍名單"):
        R.confirm_revision(ctx, rev)
    assert R.snapshot_members(ctx.con, sid) == []
    reason = next(iter(R.pending_reasons([dict(r) for r in ctx.con.execute(
        "SELECT * FROM scope_member WHERE snapshot_id=?", (sid,))])))
    assert R.bulk_member_decision(ctx, sid, reason, "納入", "人工核對同廠牌") == 3
    R.confirm_snapshot(ctx, sid)
    assert len(R.snapshot_members(ctx.con, sid)) == 3


def test_T08_scope_vs_raised_vs_annotated(ctx, tmp_path):
    rows = [[f"A0000{i:02d}100", f"D{i}", "甲", "G9", "A01AA01", "11", "錠", "N", "", "", "10", "0"] for i in range(10)]
    iid = M.stage_import(ctx, write_master(tmp_path / "t8.csv", rows), MAPPING, base_date="2026-10-01")
    M.activate(ctx, iid)
    _, rev, sid = recognized(ctx, iid, {"field": "group_code", "op": "eq", "value": "G9"})
    assert len(R.revision_scope(ctx.con, rev)) == 10
    pas = [ready_action(ctx, rev, f"A0000{i:02d}100") for i in (0, 1)]
    ans = [P.create_annotation(ctx, drug_code=f"A0000{i:02d}100", action="新增", category="特殊藥品", revision_id=rev,
                               price_action_id=pa, effective_date="2026-11-01", basis="公告")
           for i, pa in zip((0, 1), pas)]
    P.confirm_announcement(ctx, price_action_ids=pas, annotation_ids=ans)
    assert len(R.revision_scope(ctx.con, rev)) == 10
    assert ctx.con.execute("SELECT COUNT(*) FROM price_action WHERE confirm_status='已確認'").fetchone()[0] == 2
    assert ctx.con.execute("SELECT COUNT(*) FROM annotation_event WHERE confirm_status='已確認'").fetchone()[0] == 2
    assert P.annotation_on(ctx.con, "A000005100", "2026-12-01")["category"] is None
    with pytest.raises(BizError, match="範圍外"):
        ready_action(ctx, rev, "Z999999999")
    with pytest.raises(BizError, match="只有實際調高"):
        P.create_annotation(ctx, drug_code="A000005100", action="新增", category="特殊藥品", effective_date="2026-11-01")


def test_T09_contract_required_but_unsigned(ctx, master):
    _, rev, _ = recognized(ctx, master, {"field": "drug_code", "op": "eq", "value": "B000004100"},
                           category="不可替代特殊藥品")
    cid = P.create_contract(ctx, contract_no="C-9", status="辦理中", drugs=["B000004100"])
    pa = ready_action(ctx, rev, "B000004100", contract_requirement="須簽約", contract_id=cid)
    with pytest.raises(BizError, match="缺件"):
        P.confirm_announcement(ctx, price_action_ids=[pa])
    pa2 = ready_action(ctx, rev, "B000004100", contract_requirement="須簽約")
    errs = P.validate_price_action(ctx.con, ctx.con.execute("SELECT * FROM price_action WHERE id=?", (pa2,)).fetchone())
    assert any("缺件" in e for e in errs)
    assert ctx.con.execute("SELECT COUNT(*) FROM price_action WHERE confirm_status='已確認'").fetchone()[0] == 0


def test_T10_historical_no_contract_needed(ctx, master):
    _, rev, _ = recognized(ctx, master, {"field": "drug_code", "op": "eq", "value": "B000004100"})
    no_basis = ready_action(ctx, rev, "B000004100", requirement_basis=None)
    with pytest.raises(BizError, match="說明／依據"):
        P.confirm_announcement(ctx, price_action_ids=[no_basis])
    ok = ready_action(ctx, rev, "B000004100", requirement_basis="97年規定僅不可替代類別須簽約")
    P.confirm_announcement(ctx, price_action_ids=[ok])
    row = ctx.con.execute("SELECT * FROM price_action WHERE id=?", (ok,)).fetchone()
    assert row["confirm_status"] == "已確認" and row["contract_requirement"] == "當時無須簽約"


def test_T11_confirmed_but_future(ctx, master):
    rec, rev, _ = recognized(ctx, master, {"field": "drug_code", "op": "eq", "value": "B000004100"}, eff="2026-11-01")
    assert R.status_on(ctx.con, rec, "2026-10-17")["state"] == "待生效"
    assert R.status_on(ctx.con, rec, "2026-11-01")["state"] == "有效"


def test_T12_cancel_then_rerecognize(ctx, master):
    rule = {"field": "drug_code", "op": "eq", "value": "B000004100"}
    rec, rev1, sid1 = recognized(ctx, master, rule, eff="2020-01-01")
    rev2 = R.create_revision(ctx, rec, category="取消認定", effective_from="2022-01-01", completeness="完整")
    R.confirm_revision(ctx, rev2)
    rev3 = R.create_revision(ctx, rec, category="不可替代特殊藥品", rule=rule, effective_from="2024-01-01")
    sid3 = R.save_snapshot(ctx, rev3, master)
    R.confirm_snapshot(ctx, sid3)
    R.confirm_revision(ctx, rev3)
    assert R.status_on(ctx.con, rec, "2021-06-01")["state"] == "有效"
    assert R.status_on(ctx.con, rec, "2023-06-01")["state"] == "已取消"
    st = R.status_on(ctx.con, rec, "2025-06-01")
    assert st["state"] == "有效" and st["revision"]["category"] == "不可替代特殊藥品"
    r1 = ctx.con.execute("SELECT confirm_status, category FROM recognition_revision WHERE id=?", (rev1,)).fetchone()
    assert tuple(r1) == ("已確認", "特殊藥品")
    assert ctx.con.execute("SELECT status FROM scope_snapshot WHERE id=?", (sid1,)).fetchone()[0] == "已確認"
    with pytest.raises(BizError, match="已確認"):
        R.update_revision(ctx, rev1, {"category": "其他"})


def test_T13_new_master_new_match(ctx, master, tmp_path):
    _, rev, _ = recognized(ctx, master, {"field": "group_code", "op": "eq", "value": "G1CCC0000001"})
    rows = BASE_ROWS + [["B000009100", "DRUG C2", "己藥廠", "G1CCC0000001", "L04AA03", "11", "錠", "N", "", "", "5", "0"]]
    iid2 = M.stage_import(ctx, write_master(tmp_path / "m2.csv", rows), MAPPING, base_date="2026-11-01")
    M.activate(ctx, iid2)
    P.generate_diff_todos(ctx, iid2)
    todo = ctx.con.execute("SELECT * FROM review_item WHERE kind='新符合'").fetchall()
    assert [t["drug_code"] for t in todo] == ["B000009100"]
    assert "B000009100" not in R.revision_scope(ctx.con, rev)
    assert ctx.con.execute("SELECT COUNT(*) FROM price_action").fetchone()[0] == 0
    assert ctx.con.execute("SELECT COUNT(*) FROM annotation_event").fetchone()[0] == 0


def test_T14_new_master_changes_group(ctx, master, tmp_path):
    _, rev, sid = recognized(ctx, master, {"field": "drug_code", "op": "eq", "value": "B000004100"})
    rows = [list(r) for r in BASE_ROWS]
    rows[3][3] = "G1ZZZ0000001"
    iid2 = M.stage_import(ctx, write_master(tmp_path / "m3.csv", rows), MAPPING, base_date="2026-11-01")
    M.activate(ctx, iid2)
    P.generate_diff_todos(ctx, iid2)
    assert ctx.con.execute("SELECT group_code FROM scope_member WHERE snapshot_id=?", (sid,)).fetchone()[0] == "G1CCC0000001"
    assert ctx.con.execute("SELECT COUNT(*) FROM review_item WHERE kind='分組異動'").fetchone()[0] == 1
    d = M.diff_imports(ctx.con, master, iid2)
    assert d["changed"][0]["diffs"]["group_code"] == ["G1CCC0000001", "G1ZZZ0000001"]


def test_T15_contract_with_supplement(ctx, tmp_path):
    c1 = P.create_contract(ctx, contract_no="K-1", firm="甲", status="完成簽署", signed_date="2025-01-01",
                           drugs=["A000001100", "A000002100"])
    f = tmp_path / "合約.pdf"
    f.write_bytes(b"v1")
    add_attachment(ctx, f, "contract", c1, tmp_path / "pending")
    f.write_bytes(b"v2")
    add_attachment(ctx, f, "contract", c1, tmp_path / "pending")
    c2 = P.create_contract(ctx, contract_no="K-1-補1", status="完成簽署", signed_date="2026-01-01",
                           original_contract_id=c1, drugs=["A000003100"])
    chain = P.contract_chain(ctx.con, c2)
    assert [c["contract_no"] for c in chain] == ["K-1", "K-1-補1"]
    assert {d["drug_code"] for d in chain[0]["drugs"]} == {"A000001100", "A000002100"}
    assert [a["version"] for a in chain[0]["attachments"]] == [1, 2]
    keys = [a["storage_key"] for a in chain[0]["attachments"]]
    assert all((tmp_path / "pending" / k).exists() for k in keys) and keys[0] != keys[1]


# ---------------- 共用槽 ----------------

def test_T16_two_editors(tmp_path):
    s = new_store(tmp_path)
    l1 = s.acquire_lock("u1")
    with pytest.raises(LockBusy):
        s.acquire_lock("u2")
    code = ("import sys;sys.path.insert(0,%r);from specdrug.store import SharedStore,LockBusy\n"
            "try:\n SharedStore(%r).acquire_lock('u3');print('GOT')\nexcept LockBusy:\n print('BUSY')") % (
        str(ROOT), str(s.root))
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True).stdout
    assert "BUSY" in out
    s.release_lock(l1)
    s.release_lock(s.acquire_lock("u2"))


def test_T17_double_confirm_single_result(ctx, master):
    rec = R.create_recognition(ctx, name="T17")
    rev = R.create_revision(ctx, rec, category="特殊藥品", rule={"field": "drug_code", "op": "eq", "value": "B000004100"},
                            effective_from="2020-01-01")
    R.confirm_snapshot(ctx, R.save_snapshot(ctx, rev, master))
    assert R.confirm_revision(ctx, rev, op_id="op-1") == R.confirm_revision(ctx, rev, op_id="op-1")
    pa = ready_action(ctx, rev, "B000004100")
    r1 = P.confirm_announcement(ctx, price_action_ids=[pa], op_id="op-2")
    r2 = P.confirm_announcement(ctx, price_action_ids=[pa], op_id="op-2")
    assert r1 == r2
    n = ctx.con.execute("SELECT COUNT(*) FROM audit_event WHERE entity='price_action' AND entity_id=? AND diff_json LIKE"
                        " '%已確認%'", (pa,)).fetchone()[0]
    assert n == 1


def test_T18_failure_no_partial_publish(ctx, master, tmp_path):
    _, rev, _ = recognized(ctx, master, {"field": "group_code", "op": "eq", "value": "G1AAA0000001"}, decide_pending="納入")
    good = ready_action(ctx, rev, "A000001100")
    bad = ready_action(ctx, rev, "A000002100", announce_doc_no=None)
    with pytest.raises(BizError, match="整批未確認"):
        P.confirm_announcement(ctx, price_action_ids=[good, bad])
    assert ctx.con.execute("SELECT COUNT(*) FROM price_action WHERE confirm_status='已確認'").fetchone()[0] == 0
    # 附件不存在→整體不發布、指標不變
    s = new_store(tmp_path)
    ws = Workspace(tmp_path / "loc", s.root, "admin")
    ws.refresh_view()
    ws.start_edit()
    before = s.current()
    f = tmp_path / "x.pdf"
    f.write_bytes(b"abc")
    cid = C.create_case(ws.ctx(), subject="附件案")
    add_attachment(ws.ctx(), f, "case_record", cid, ws.pending_dir)
    for p in ws.pending_dir.rglob("*"):
        if p.is_file():
            p.unlink()
    with pytest.raises(StoreError, match="附件"):
        ws.publish()
    assert s.current() == before
    assert not any(s.staging.iterdir()) or all(not p.name.endswith(".partial") for p in s.staging.iterdir())
    ws.close()


def test_T19_category_conflict(ctx, master):
    rule = {"field": "drug_code", "op": "eq", "value": "B000004100"}
    recognized(ctx, master, rule, name="甲認定", category="特殊藥品")
    rec2 = R.create_recognition(ctx, name="乙認定")
    rev2 = R.create_revision(ctx, rec2, category="不可替代特殊藥品", rule=rule, effective_from="2021-01-01")
    R.confirm_snapshot(ctx, R.save_snapshot(ctx, rev2, master))
    with pytest.raises(BizError, match="衝突"):
        R.confirm_revision(ctx, rev2)
    R.confirm_revision(ctx, rev2, conflict_basis="擬訂會議決議另案處理")
    r = R.drug_category_on(ctx.con, "B000004100", "2026-10-17")
    assert r["conflict"] is True and r["category"] is None and len(r["sources"]) == 2


def test_T20_contract_expiry_no_auto_cancel(ctx, master):
    _, rev, _ = recognized(ctx, master, {"field": "drug_code", "op": "eq", "value": "B000004100"})
    cid = signed_contract(ctx, ["B000004100"], expiry="2026-10-01")
    pa = ready_action(ctx, rev, "B000004100", contract_requirement="須簽約", contract_id=cid,
                      effective_date="2026-09-01")
    an = P.create_annotation(ctx, drug_code="B000004100", action="新增", category="特殊藥品", revision_id=rev,
                             price_action_id=pa, effective_date="2026-09-01", basis="公告")
    P.confirm_announcement(ctx, price_action_ids=[pa], annotation_ids=[an])
    alerts = P.contract_alerts(ctx.con, "2026-10-17")
    assert alerts and alerts[0]["alert"] == "已到期"
    assert R.drug_category_on(ctx.con, "B000004100", "2026-10-17")["category"] == "特殊藥品"
    assert P.annotation_on(ctx.con, "B000004100", "2026-10-17")["category"] == "特殊藥品"
    assert ctx.con.execute("SELECT confirm_status FROM price_action WHERE id=?", (pa,)).fetchone()[0] == "已確認"


def test_T21_drug_timeline(ctx, master):
    cid = C.create_case(ctx, subject="申請B", recv_date="2026-01-05", drugs=["B000004100"])
    cd = ctx.con.execute("SELECT id FROM case_drug WHERE case_id=?", (cid,)).fetchone()[0]
    mid = C.create_meeting(ctx, meeting_type="共同擬訂", name="擬訂會議115年6月", meeting_date="2026-06-18")
    ag = C.add_agenda(ctx, meeting_id=mid, case_id=cid)
    C.add_decision(ctx, agenda_id=ag, nature="會議決議", result="同意", drug_results={cd: "同意"})
    _, rev, _ = recognized(ctx, master, {"field": "drug_code", "op": "eq", "value": "B000004100"})
    cont = signed_contract(ctx, ["B000004100"])
    pa = ready_action(ctx, rev, "B000004100", case_id=cid, contract_requirement="須簽約", contract_id=cont)
    an = P.create_annotation(ctx, drug_code="B000004100", action="新增", category="特殊藥品", revision_id=rev,
                             price_action_id=pa, effective_date="2026-11-01", basis="公告")
    P.confirm_announcement(ctx, price_action_ids=[pa], annotation_ids=[an])
    d = drug_detail(ctx.con, "B000004100", "2026-12-01")
    kinds = {e["kind"] for e in d["events"]}
    assert {"收案", "共同擬訂會議／會議決議", "認定", "調高", "合約", "註記"} <= kinds
    assert d["annotation"]["category"] == "特殊藥品" and d["master"]


def test_T22_backup_restore(tmp_path):
    s = new_store(tmp_path)
    ws = Workspace(tmp_path / "loc", s.root, "admin")
    ws.refresh_view()
    ws.start_edit()
    ctx = ws.ctx()
    cid = C.create_case(ctx, subject="備份案", drugs=["A000001100"])
    f = tmp_path / "doc.pdf"
    f.write_bytes(b"content")
    add_attachment(ctx, f, "case_record", cid, ws.pending_dir)
    rid1 = ws.publish()
    bk = s.backup(rid1)
    assert verify_backup(bk) == []
    C.create_case(ws.ctx(), subject="之後的案")
    ws.publish()
    rid3 = s.restore(ws.lock, bk, operator="admin", reason="演練", op_id="restore-1")
    assert s.current()["release_id"] == rid3 and s.verify_release(rid3)
    assert (s.releases / rid1).exists()  # 歷史release未被改寫
    ws.end_edit("discard")
    con = ws.con
    assert [r[0] for r in con.execute("SELECT subject FROM case_record")] == ["備份案"]
    a = con.execute("SELECT * FROM attachment").fetchone()
    assert con.execute("SELECT entity_id FROM attachment_link WHERE attachment_id=?", (a["id"],)).fetchone()[0] == cid
    out = s.fetch_attachment(a["storage_key"], a["sha256"], tmp_path / "dl.pdf")
    assert out.read_bytes() == b"content"
    assert con.execute("SELECT COUNT(*) FROM audit_event WHERE action='還原'").fetchone()[0] == 1
    # 竄改備份→驗證失敗
    (bk / "data.sqlite").write_bytes(b"x")
    assert verify_backup(bk)
    ws.close()


def test_T23_viewer_cannot_modify(ctx, tmp_path):
    viewer = Ctx(ctx.con, "viewer", {ROLE_VIEWER}, True)
    with pytest.raises(PermissionDenied):
        C.create_case(viewer, subject="x")
    s = new_store(tmp_path)
    ws = Workspace(tmp_path / "loc", s.root, "nobody")
    ws.refresh_view()
    with pytest.raises(PermissionDenied, match="查詢模式"):
        C.create_case(ws.ctx(), subject="x")
    with pytest.raises(StoreError, match="未授權"):
        ws.start_edit()
    assert not ws.editing
    with pytest.raises(Exception):
        ws.con.execute("INSERT INTO meta(key,value) VALUES ('x','y')")  # 本機副本唯讀
    ws.close()
    # 共用槽目錄寫入權限須由機關檔案伺服器設定（見 docs/驗收結果.md，人工驗證）


def test_T24_rufinamide_not_recognized(ctx, tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "final"
    ws.append([MG.COLS[k] for k in MG.COLS])
    row = {k: None for k in MG.COLS}
    row.update(seq="172", group_code="N03RUF212116", category="特殊藥品", rule="分組代碼=N03RUF212116",
               price_history="Rufinamide參考特殊藥品精神調高，非認列特殊藥品", history="擬訂會議113年8月")
    ws.append([row[k] for k in MG.COLS])
    wb.save(tmp_path / "w.xlsx")
    MG.stage_workbook(ctx, tmp_path / "w.xlsx")
    r = MG.list_rows(ctx.con)[0]
    res = MG.commit_row(ctx, r["id"])
    assert res.get("recognition_id") is None
    assert ctx.con.execute("SELECT COUNT(*) FROM recognition").fetchone()[0] == 0
    case = ctx.con.execute("SELECT notes FROM case_record WHERE id=?", (res["case_id"],)).fetchone()[0]
    assert "非認列" in case


def test_T25_unpublished_edit_invisible(tmp_path):
    s = new_store(tmp_path)
    ed = Workspace(tmp_path / "ed", s.root, "admin")
    ed.refresh_view()
    ed.start_edit()
    C.create_case(ed.ctx(), subject="草稿案")
    q = Workspace(tmp_path / "q", s.root, "viewer")
    q.refresh_view()
    assert q.con.execute("SELECT COUNT(*) FROM case_record").fetchone()[0] == 0
    assert "目前編輯者：admin" in q.banner()
    ed.publish()
    assert q.check_newer()
    q.refresh_view()
    assert q.con.execute("SELECT COUNT(*) FROM case_record").fetchone()[0] == 1
    ed.close()
    q.close()


def test_T26_interrupted_upload(tmp_path, monkeypatch):
    s = new_store(tmp_path)
    ws = Workspace(tmp_path / "loc", s.root, "admin")
    ws.refresh_view()
    ws.start_edit()
    C.create_case(ws.ctx(), subject="斷線案")
    before = s.current()

    def boom(*a, **k):
        raise OSError("網路中斷")
    monkeypatch.setattr("specdrug.store.os.rename", boom)
    with pytest.raises(OSError):
        ws.publish()
    assert s.current() == before
    assert not [p for p in s.staging.iterdir() if p.name.endswith(".partial")]
    q = Workspace(tmp_path / "q", s.root, "v")
    q.refresh_view()
    assert q.view_release == before["release_id"]
    ws.close()
    q.close()


def test_T27_crash_recovery(tmp_path):
    s = new_store(tmp_path)
    code = ("import sys,time;sys.path.insert(0,%r);from specdrug.workspace import Workspace\n"
            "from specdrug import cases as C\nw=Workspace(%r,%r,'admin');w.refresh_view();w.start_edit()\n"
            "C.create_case(w.ctx(),subject='當機前草稿');print('READY',flush=True);time.sleep(60)") % (
        str(ROOT), str(tmp_path / "ed"), str(s.root))
    p = subprocess.Popen([sys.executable, "-c", code], stdout=subprocess.PIPE, text=True)
    assert p.stdout.readline().strip() == "READY"
    with pytest.raises(LockBusy):
        s.acquire_lock("other")
    p.kill()
    p.wait()
    ws = Workspace(tmp_path / "ed", s.root, "admin")
    ws.refresh_view()
    assert ws.start_edit() == "resumed"  # 編輯權由檔案伺服端（OS）於程序結束釋放，非逾時刪檔
    assert ws.con.execute("SELECT subject FROM case_record").fetchone()[0] == "當機前草稿"
    assert (s.control / "edit.lock").exists()
    ws.close()


def test_T28_old_editor_revives(tmp_path):
    s = new_store(tmp_path)
    old = Workspace(tmp_path / "old", s.root, "admin")
    old.refresh_view()
    old.start_edit()
    C.create_case(old.ctx(), subject="舊編輯者")
    old_lock = old.lock
    old_lock.release()  # 模擬斷線：伺服端已釋放其鎖
    assert not old.lock_valid()
    new = Workspace(tmp_path / "new", s.root, "admin")
    new.refresh_view()
    new.start_edit()
    C.create_case(new.ctx(), subject="新編輯者")
    new.publish()
    with pytest.raises(LockLost):
        old.publish()
    # 重新取得編輯權後也不得整份覆蓋
    new.end_edit("keep")
    old.lock = s.acquire_lock("admin")
    with pytest.raises(PublishConflict):
        old.publish()
    old.close()
    assert old.start_edit() == "recovered"
    assert [r[0] for r in old.con.execute("SELECT subject FROM case_record")] == ["新編輯者"]
    old.close()


def test_T29_retry_after_pointer_switch_failure(tmp_path, monkeypatch):
    s = new_store(tmp_path)
    ws = Workspace(tmp_path / "loc", s.root, "admin")
    ws.refresh_view()
    ws.start_edit()
    C.create_case(ws.ctx(), subject="T29")
    orig = SharedStore._switch
    calls = {"n": 0}

    def flaky(self, lock, manifest):
        calls["n"] += 1
        if calls["n"] == 1:
            raise OSError("切換時斷線")
        return orig(self, lock, manifest)
    monkeypatch.setattr(SharedStore, "_switch", flaky)
    base = s.current()["release_id"]
    with pytest.raises(OSError):
        ws.publish()
    assert s.current()["release_id"] == base
    rid = ws.publish()
    ops = [m for m in s.list_releases() if m["parent"] == base]
    assert len(ops) == 1 and ops[0]["release_id"] == rid == s.current()["release_id"]
    ws.close()


def test_T30_consistent_view_during_publish(tmp_path):
    s = new_store(tmp_path)
    q = Workspace(tmp_path / "q", s.root, "v")
    q.refresh_view()
    first = q.view_release
    ed = Workspace(tmp_path / "ed", s.root, "admin")
    ed.refresh_view()
    ed.start_edit()
    C.create_case(ed.ctx(), subject="新版")
    ed.publish()
    assert q.view_release == first
    assert q.con.execute("SELECT value FROM meta WHERE key='release_id'").fetchone()[0] == first
    assert q.check_newer() and "新版" in q.banner()
    ed.close()
    q.close()


def test_T31_same_account_two_machines(tmp_path):
    s = new_store(tmp_path)
    a = Workspace(tmp_path / "pc1", s.root, "admin")
    b = Workspace(tmp_path / "pc2", s.root, "admin")
    a.refresh_view()
    b.refresh_view()
    a.start_edit()
    with pytest.raises(LockBusy):
        b.start_edit()
    a.close()
    b.close()


def test_T32_unreliable_share_blocks_edit(tmp_path, monkeypatch):
    from specdrug import store as S
    monkeypatch.setattr(S.EditLock, "acquire", lambda self: setattr(self, "fh", open(self.path, "a+b")))
    monkeypatch.setattr(S, "probe_lock", None)
    s = SharedStore(tmp_path / "bad")
    with pytest.raises(StoreError, match="不支援可靠鎖定"):
        s.init("admin")
    assert not s.environment_ok()
    ws = Workspace(tmp_path / "loc", None, "admin")  # 僅能離線原型
    assert ws.mode == "offline" and "離線原型" in ws.banner()
    ws.close()


# ---------------- 收案與會議 ----------------

def test_T33_rejected_case_with_two_meetings(ctx):
    cid = C.create_case(ctx, subject="未通過案", applicant="甲", drugs=["A000001100"])
    cd = ctx.con.execute("SELECT id FROM case_drug WHERE case_id=?", (cid,)).fetchone()[0]
    m1 = C.create_meeting(ctx, meeting_type="專家", name="專家會議115年3月", meeting_date="2026-03-10")
    m2 = C.create_meeting(ctx, meeting_type="共同擬訂", name="擬訂會議115年4月", meeting_date="2026-04-16")
    d1 = C.add_decision(ctx, agenda_id=C.add_agenda(ctx, meeting_id=m1, case_id=cid), nature="專家建議",
                        result="不同意", drug_results={cd: "不同意"})
    d2 = C.add_decision(ctx, agenda_id=C.add_agenda(ctx, meeting_id=m2, case_id=cid), nature="會議決議",
                        result="不同意", drug_results={cd: "不同意"})
    C.set_case_summary(ctx, cid, "全部不同意", [d1, d2])
    C.close_case(ctx, cid, status="已結案", reason="共擬不同意", basis="擬訂會議115年4月紀錄")
    tl = C.case_timeline(ctx.con, cid)
    assert [a["meeting_type"] for a in tl] == ["專家", "共同擬訂"]
    assert ctx.con.execute("SELECT COUNT(*) FROM recognition").fetchone()[0] == 0
    assert C.list_cases(ctx.con, summary="全部不同意")[0]["id"] == cid


def test_T34_expert_resubmit_joint_reject(ctx):
    cid = C.create_case(ctx, subject="T34", drugs=["A000001100"])
    cd = ctx.con.execute("SELECT id FROM case_drug WHERE case_id=?", (cid,)).fetchone()[0]
    m1 = C.create_meeting(ctx, meeting_type="專家", name="專家1", meeting_date="2026-01-01")
    m2 = C.create_meeting(ctx, meeting_type="專家", name="專家2", meeting_date="2026-03-01")
    m3 = C.create_meeting(ctx, meeting_type="共同擬訂", name="共擬", meeting_date="2026-05-01")
    C.add_decision(ctx, agenda_id=C.add_agenda(ctx, meeting_id=m1, case_id=cid), nature="專家建議",
                   result="補充資料後再議", drug_results={cd: "補充資料後再議"})
    C.add_decision(ctx, agenda_id=C.add_agenda(ctx, meeting_id=m2, case_id=cid), nature="專家建議", result="同意",
                   drug_results={cd: "同意"})
    d3 = C.add_decision(ctx, agenda_id=C.add_agenda(ctx, meeting_id=m3, case_id=cid), nature="會議決議",
                        result="不同意", drug_results={cd: "不同意"})
    tl = C.case_timeline(ctx.con, cid)
    assert [a["decisions"][0]["result"] for a in tl] == ["補充資料後再議", "同意", "不同意"]
    assert C.drug_followup_basis(ctx.con, cid, "A000001100") == []  # 專家同意不當最終通過
    with pytest.raises(BizError, match="依據"):
        C.set_case_summary(ctx, cid, "全部不同意", [])
    C.set_case_summary(ctx, cid, "全部不同意", [d3])
    assert ctx.con.execute("SELECT result_summary FROM case_record WHERE id=?", (cid,)).fetchone()[0] == "全部不同意"


def test_T35_three_drugs_different_results(ctx, master):
    cid = C.create_case(ctx, subject="T35", drugs=["A000001100", "A000002100", "A000006100"])
    cds = {r["drug_code"]: r["id"] for r in ctx.con.execute("SELECT * FROM case_drug WHERE case_id=?", (cid,))}
    m = C.create_meeting(ctx, meeting_type="共同擬訂", name="共擬T35", meeting_date="2026-06-01")
    C.add_decision(ctx, agenda_id=C.add_agenda(ctx, meeting_id=m, case_id=cid), nature="會議決議", result="其他",
                   result_text="逐品項不同", drug_results={cds["A000001100"]: "同意", cds["A000002100"]: "不同意",
                                                          cds["A000006100"]: "暫緩"})
    _, rev, _ = recognized(ctx, master, {"field": "group_code", "op": "eq", "value": "G1AAA0000001"},
                           decide_pending="納入")
    assert ready_action(ctx, rev, "A000001100", case_id=cid)
    for code in ("A000002100", "A000006100"):
        with pytest.raises(BizError, match="同意／附條件同意"):
            ready_action(ctx, rev, code, case_id=cid)


def test_T36_one_meeting_many_cases(ctx):
    m = C.create_meeting(ctx, meeting_type="共同擬訂", name="擬訂會議115年6月", meeting_date="2026-06-18")
    c1 = C.create_case(ctx, subject="甲案", drugs=["A000001100"])
    c2 = C.create_case(ctx, subject="乙案", drugs=["A000002100"])
    a1 = C.add_agenda(ctx, meeting_id=m, case_id=c1, agenda_no="1")
    a2 = C.add_agenda(ctx, meeting_id=m, case_id=c2, agenda_no="2")
    C.add_decision(ctx, agenda_id=a1, nature="會議決議", result="同意")
    C.add_decision(ctx, agenda_id=a2, nature="會議決議", result="不同意")
    assert C.case_timeline(ctx.con, c1)[0]["decisions"][0]["result"] == "同意"
    assert C.case_timeline(ctx.con, c2)[0]["decisions"][0]["result"] == "不同意"
    cd2 = ctx.con.execute("SELECT id FROM case_drug WHERE case_id=?", (c2,)).fetchone()[0]
    d1 = ctx.con.execute("SELECT id FROM decision WHERE agenda_id=?", (a1,)).fetchone()[0]
    with pytest.raises(BizError, match="不能串接"):
        C.set_decision_drug(ctx, d1, cd2, "同意")
    assert ctx.con.execute("SELECT COUNT(*) FROM meeting").fetchone()[0] == 1


@pytest.mark.skipif(not UPLOAD.exists(), reason="需使用者提供之活頁簿1.xlsx")
def test_T37_T38_real_workbook(ctx):
    r1 = MG.stage_workbook(ctx, UPLOAD)
    rows = MG.list_rows(ctx.con)
    for r in rows:
        MG.commit_row(ctx, r["id"])
    assert ctx.con.execute("SELECT COUNT(*) FROM recognition_revision WHERE confirm_status<>'草稿'").fetchone()[0] == 0
    assert ctx.con.execute("SELECT COUNT(*) FROM annotation_event").fetchone()[0] == 0
    cancel = ctx.con.execute("SELECT COUNT(*) FROM recognition_revision WHERE category='取消認定'").fetchone()[0]
    assert cancel == 4  # 5列取消註記，其中Rufinamide(第172列)非認列不建認定
    no_date = ctx.con.execute("SELECT COUNT(*) FROM recognition_revision WHERE effective_from IS NULL AND"
                              " completeness='不完整'").fetchone()[0]
    assert no_date > 200
    for rec in ctx.con.execute("SELECT id FROM recognition"):
        assert R.status_on(ctx.con, rec[0])["state"] in ("草稿／待確認", "無已確認版本")
    # T38：同一Excel重複匯入
    n_case = ctx.con.execute("SELECT COUNT(*) FROM case_record").fetchone()[0]
    n_meet = ctx.con.execute("SELECT COUNT(*) FROM meeting").fetchone()[0]
    r2 = MG.stage_workbook(ctx, UPLOAD)
    assert r2["new"] == 0 and r2["already_staged"] == r1["total"]
    with pytest.raises(BizError, match="不重複建案"):
        MG.commit_row(ctx, rows[0]["id"])
    assert ctx.con.execute("SELECT COUNT(*) FROM case_record").fetchone()[0] == n_case
    assert ctx.con.execute("SELECT COUNT(*) FROM meeting").fetchone()[0] == n_meet


def test_T37_synthetic_missing_date_and_cancel(ctx, tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "final"
    ws.append([MG.COLS[k] for k in MG.COLS])
    base = {k: None for k in MG.COLS}
    r1 = dict(base, seq="1", group_code="L01DOX412202", cancel="取消", category="取消特殊藥品註記",
              rule="分組代碼=L01DOX412202", history="擬訂會議110年6月")
    r2 = dict(base, seq="2", group_code="X01", category="不可替代必要藥品(104.05.01藥價歸零)",
              rule="分組代碼=X01", history="藥事小組第9-11次 (101年2月9日)\n擬訂會議113年8月")
    for r in (r1, r2):
        ws.append([r[k] for k in MG.COLS])
    wb.save(tmp_path / "s.xlsx")
    MG.stage_workbook(ctx, tmp_path / "s.xlsx")
    rows = MG.list_rows(ctx.con)
    res1 = MG.commit_row(ctx, rows[0]["id"])
    res2 = MG.commit_row(ctx, rows[1]["id"])
    assert "cancel_revision_id" in res1
    rv = ctx.con.execute("SELECT * FROM recognition_revision WHERE id=?", (res2["revision_id"],)).fetchone()
    assert rv["category"] == "不可替代特殊藥品" and rv["original_category"].startswith("不可替代必要藥品")
    assert rv["effective_from"] is None and rv["completeness"] == "不完整" and rv["confirm_status"] == "草稿"
    assert any("藥價歸零" in f for f in rows[1]["parsed"]["category_flags"])
    tl = C.case_timeline(ctx.con, res2["case_id"])
    assert [a["meeting_type"] for a in tl] == ["藥事小組", "共同擬訂"]  # 拆成兩個議案，不合併成最後一次
    assert tl[0]["meeting_date"] == "2012-02-09" and tl[1]["meeting_date"] is None  # 只有年月者不補日


def test_T39_rejection_does_not_cancel_existing(ctx, master):
    rec, rev, _ = recognized(ctx, master, {"field": "drug_code", "op": "eq", "value": "B000004100"})
    cid = C.create_case(ctx, subject="新申請", drugs=["B000004100"])
    cd = ctx.con.execute("SELECT id FROM case_drug WHERE case_id=?", (cid,)).fetchone()[0]
    m = C.create_meeting(ctx, meeting_type="共同擬訂", name="共擬T39", meeting_date="2026-08-01")
    d = C.add_decision(ctx, agenda_id=C.add_agenda(ctx, meeting_id=m, case_id=cid), nature="會議決議",
                       result="不同意", drug_results={cd: "不同意"})
    C.set_case_summary(ctx, cid, "全部不同意", [d])
    assert R.status_on(ctx.con, rec, "2026-10-17")["state"] == "有效"
    assert R.drug_category_on(ctx.con, "B000004100", "2026-10-17")["category"] == "特殊藥品"


def test_T40_no_fabricated_cases(ctx, tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "final"
    ws.append([MG.COLS[k] for k in MG.COLS])
    ws.append(["1", None, "G1", None, None, None, None, "擬訂會議113年8月", None, "特殊藥品", None, None,
               "分組代碼=G1", None, None, None, None])
    wb.save(tmp_path / "one.xlsx")
    MG.stage_workbook(ctx, tmp_path / "one.xlsx")
    for r in MG.list_rows(ctx.con):
        MG.commit_row(ctx, r["id"])
    assert ctx.con.execute("SELECT COUNT(*) FROM case_record").fetchone()[0] == 1
    assert ctx.con.execute("SELECT COUNT(*) FROM case_record WHERE result_summary<>'尚無結論'").fetchone()[0] == 0
    # 人工補登入口
    cid = C.create_case(ctx, subject="人工補登之未通過案", applicant="乙")
    assert C.list_cases(ctx.con, keyword="人工補登")[0]["id"] == cid
