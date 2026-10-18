"""既有Excel（活頁簿1 final）歷史資料移轉：暫存、預覽、逐列確認寫入。

- 每列保存來源雜湊、工作表、列號、原始欄位、解析結果、待確認問題。
- 不直接寫成現行有效資料：認定版本一律為草稿；缺日期／決議／合約標記不完整。
- 同來源同列重複匯入不重複建案；相似列僅列合併候選，不自行合併。
- 未提供的未通過案件不能編造。
"""
from __future__ import annotations

import json
import os
import re
from collections import defaultdict

from . import rules
from .cases import create_case, create_meeting, find_meeting, add_agenda, add_decision
from .core import BizError, Ctx, ROLE_CLERK, ROLE_REVIEWER, audit, get
from .recognition import create_recognition, create_revision
from .util import new_id, parse_date, sha256_file, utc_now

COLS = {
    "seq": "調整時公告序號", "cancel": "取消註記", "group_code": "分組代碼", "group_name": "分組名稱",
    "group_item": "分組項目", "ingredient": "成分/含量", "dosage": "劑型", "history": "討論歷程",
    "price_history": "價格異動歷程", "category": "類別", "side_decision": "附帶決議", "ann_name": "調整時公告名稱",
    "rule": "篩選原則", "review_10512": "原特殊品項_10512擬訂會議檢討結果", "old_rule": "原篩選條件",
    "eff_date": "特殊藥品生效日期", "contract_codes": "有簽訂協議為不可替代必要/特殊(藥品代碼)",
}
HIST_TO_STD = {"不可替代必要藥品": "不可替代特殊藥品", "必要藥品": "特殊藥品",
               "不可替代特殊藥品": "不可替代特殊藥品", "特殊藥品": "特殊藥品"}
_MEET_DATE = re.compile(r"(\d{2,3})年(\d{1,2})月(?:(\d{1,2})日)?")
_CODE = re.compile(r"\b[A-Z][A-Z0-9]\d{6,8}[0-9A-Z]{0,2}\b")


def _norm(h):
    return re.sub(r"\s+", "", str(h or ""))


def _match_cols(headers):
    idx = {}
    for key, name in COLS.items():
        for i, h in enumerate(headers):
            if _norm(h).startswith(_norm(name)):
                idx[key] = i
                break
    missing = [COLS[k] for k in ("seq", "group_code", "category", "rule") if k not in idx]
    if missing:
        raise BizError("找不到必要欄位：" + "、".join(missing) + "；請確認是 final 工作表。")
    return idx


def _s(v):
    if v is None:
        return None
    t = str(v).strip()
    return t or None


def parse_category(raw):
    """回傳(標準類別, 旗標, 問題)。混合文字須人工拆解。"""
    flags, issues = [], []
    if not raw:
        return "待確認", flags, ["類別空白"]
    t = raw.replace("\n", "/")
    for m in re.finditer(r"(\d{2,3}\.\d{2}\.\d{2})藥價歸零", t):
        flags.append(f"藥價歸零 {m.group(1)}")
    if "未簽約" in t:
        flags.append("未簽約")
    if "取消" in t or "(刪)" in t:
        flags.append("取消候選")
    base = re.sub(r"\(.*?\)", "", t).strip(" /")
    parts = [p.strip() for p in re.split(r"[/→]", base) if p.strip()]
    if len(parts) > 1:
        issues.append(f"類別含多個或變更文字「{raw}」，須人工拆解")
        return "待確認", flags, issues
    if not parts:
        return "待確認", flags, issues + [f"類別「{raw}」無法對應"]
    std = HIST_TO_STD.get(parts[0])
    if parts[0].startswith("取消"):
        return "取消認定", flags, issues
    if re.search(r"\(取消[^)]*\)", t):
        issues.append(f"類別「{raw}」含取消／改類文字，須人工拆解")
        return "待確認", flags, issues
    if not std:
        return "待確認", flags, issues + [f"類別「{raw}」無法對應"]
    return std, flags, issues


def parse_meetings(text):
    """討論歷程拆成議案候選；每行一個，保留原文；日期只取有明確年月日者。"""
    out = []
    for line in str(text or "").split("\n"):
        ln = line.strip().rstrip("，,")
        ln = re.sub(r"^\d+\.\s*", "", ln)
        if not ln:
            continue
        if "藥事小組" in ln:
            typ = "藥事小組"
        elif "擬訂" in ln or "擬定" in ln or "共擬" in ln:
            typ = "共同擬訂"
        elif "專家" in ln:
            typ = "專家"
        else:
            typ = "其他"
        date = None
        m = _MEET_DATE.search(ln)
        if m and m.group(3):
            date = parse_date(f"{m.group(1)}/{m.group(2)}/{m.group(3)}")
        out.append({"type": typ, "name": ln, "date": date, "original": line.strip()})
    return out


def parse_row(raw: dict) -> tuple[dict, list[str]]:
    issues = []
    p = {k: _s(raw.get(k)) for k in COLS}
    std, flags, ci = parse_category(p["category"])
    issues += ci
    p["category_std"] = std
    p["category_flags"] = flags
    ast, ri = rules.parse_rule_text(p["rule"])
    issues += [f"篩選原則：{x}" for x in ri]
    p["rule_ast"] = ast
    if p["eff_date"]:
        d = parse_date(p["eff_date"])
        p["eff_date_iso"] = d
        if not d:
            issues.append(f"特殊藥品生效日期「{p['eff_date']}」無法解析")
    else:
        p["eff_date_iso"] = None
        issues.append("特殊藥品生效日期空白（未知，不以會議日補）")
    p["meetings"] = parse_meetings(p["history"])
    if not p["meetings"]:
        issues.append("無可辨識的會議紀錄")
    for m in p["meetings"]:
        if m["type"] == "其他":
            issues.append(f"討論歷程「{m['name']}」無法判斷會議類型")
        if not m["date"]:
            pass
    p["contract_code_clues"] = sorted(set(_CODE.findall(p["contract_codes"] or "")))
    if p["contract_code_clues"]:
        issues.append("有協議藥品代碼線索，僅為合約關聯線索，不視為已核驗正式合約")
    if p["cancel"]:
        issues.append("取消註記：建立取消候選版本，須補依據與日期")
    whole = " ".join(str(v) for v in raw.values() if v)
    p["not_recognized"] = bool(re.search(r"非認列|參考特殊藥品精神", whole))
    if p["not_recognized"]:
        issues.append("原文註明非認列（參考特殊藥品精神調高），保留價格案件及原文，不建立有效特殊身分")
    if p["seq"] and re.search(r"重複|包含在", p["seq"]):
        issues.append(f"公告序號註記「{p['seq']}」，可能與其他列重複，請核對")
    if p["group_item"] is not None:
        issues.append("分組項目意義未確認，僅保存原值")
    return p, issues


# ---------------- 暫存 ----------------

def stage_workbook(ctx: Ctx, path, sheet="final") -> dict:
    ctx.require(ROLE_CLERK, ROLE_REVIEWER)
    import openpyxl
    sha = sha256_file(path)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise BizError(f"找不到工作表「{sheet}」。")
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    headers = list(next(it))
    idx = _match_cols(headers)
    rows = []
    for no, r in enumerate(it, start=2):
        if r is None or all(c is None for c in r):
            continue
        rows.append((no, {k: (r[i] if i < len(r) else None) for k, i in idx.items()}))
    wb.close()
    new, existing = 0, 0
    with ctx.tx():
        batch_id = new_id()
        ctx.con.execute("INSERT INTO migration_batch(id, source_name, sha256, sheet, imported_by, imported_at, row_count)"
                        " VALUES (?,?,?,?,?,?,?)",
                        (batch_id, os.path.basename(str(path)), sha, sheet, ctx.actor, utc_now(), len(rows)))
        for no, raw in rows:
            if ctx.con.execute("SELECT 1 FROM migration_row WHERE source_sha256=? AND sheet=? AND row_no=?",
                               (sha, sheet, no)).fetchone():
                existing += 1
                continue
            parsed, issues = parse_row(raw)
            ctx.con.execute(
                "INSERT INTO migration_row(id, batch_id, source_sha256, sheet, row_no, raw_json, parsed_json, issues_json)"
                " VALUES (?,?,?,?,?,?,?,?)",
                (new_id(), batch_id, sha, sheet, no, json.dumps(raw, ensure_ascii=False, default=str),
                 json.dumps(parsed, ensure_ascii=False, default=str), json.dumps(issues, ensure_ascii=False)))
            new += 1
        audit(ctx, "migration_batch", batch_id, "暫存", after={"source": os.path.basename(str(path)), "sha256": sha,
                                                              "new": new, "existing": existing})
    return {"batch_id": batch_id, "sha256": sha, "new": new, "already_staged": existing, "total": len(rows)}


def merge_candidates(con, sha) -> list[dict]:
    """同公告序號、公告名稱或分組代碼的列，僅列為合併候選。"""
    groups = defaultdict(list)
    for r in con.execute("SELECT id, row_no, parsed_json FROM migration_row WHERE source_sha256=?", (sha,)):
        p = json.loads(r["parsed_json"])
        for key in ("seq", "ann_name", "group_code"):
            if p.get(key):
                groups[(key, p[key])].append(r["row_no"])
    return [{"key": COLS[k], "value": v, "rows": sorted(set(rows))} for (k, v), rows in groups.items()
            if len(set(rows)) > 1]


def list_rows(con, sha=None, status=None):
    sql = "SELECT * FROM migration_row WHERE 1=1"
    args = []
    if sha:
        sql += " AND source_sha256=?"
        args.append(sha)
    if status:
        sql += " AND status=?"
        args.append(status)
    sql += " ORDER BY source_sha256, row_no"
    return [dict(r, parsed=json.loads(r["parsed_json"]), issues=json.loads(r["issues_json"]))
            for r in con.execute(sql, args)]


# ---------------- 寫入 ----------------

def commit_row(ctx: Ctx, row_id, *, merge_into_recognition_id=None, rule_override=None) -> dict:
    """使用者確認後寫入案件、會議、議案（決議待確認）及認定草稿版本。"""
    ctx.require(ROLE_CLERK, ROLE_REVIEWER)
    row = get(ctx, "migration_row", row_id)
    if not row:
        raise BizError("找不到暫存列。")
    if row["status"] == "已寫入":
        raise BizError(f"第{row['row_no']}列已寫入（案件 {row['committed_case_id']}），不重複建案。")
    if row["status"] == "略過":
        raise BizError("此列已標示略過。")
    p = json.loads(row["parsed_json"])
    issues = json.loads(row["issues_json"])
    batch = get(ctx, "migration_batch", row["batch_id"])
    src = f"{batch['source_name']}#{row['sheet']}!{row['row_no']}（SHA-256 {row['source_sha256'][:12]}）"
    title = p.get("ann_name") or p.get("group_name") or p.get("group_code") or f"第{row['row_no']}列"
    notes = []
    if p.get("price_history"):
        notes.append(f"價格異動歷程（原文，待人工拆成決議／價格事件）：\n{p['price_history']}")
    if p.get("contract_code_clues"):
        notes.append("協議藥品代碼線索（未核驗）：" + "、".join(p["contract_code_clues"]))
    if p.get("review_10512"):
        notes.append(f"原特殊品項_10512擬訂會議檢討結果（歷史證據）：{p['review_10512']}")
    if p.get("group_item"):
        notes.append(f"分組項目原值：{p['group_item']}")
    if issues:
        notes.append("移轉待確認：" + "；".join(issues))
    result = {}
    with ctx.tx():
        case_id = create_case(
            ctx, subject=title, request_item="特殊藥品認定／調高（歷史移轉）",
            case_no=f"HIST-{row['source_sha256'][:6]}-{row['row_no']:04d}", completeness="不完整",
            pending_fields="recv_date,recv_doc_no,applicant", source_ref=src, notes="\n\n".join(notes))
        result["case_id"] = case_id
        for i, m in enumerate(p.get("meetings", []), start=1):
            mid = find_meeting(ctx.con, m["type"], m["name"], m["date"]) or create_meeting(
                ctx, meeting_type=m["type"], name=m["name"], meeting_date=m["date"],
                notes="歷史移轉建立；日期僅有年月者留空待確認" if not m["date"] else "歷史移轉建立")
            ag = add_agenda(ctx, meeting_id=mid, case_id=case_id, agenda_no=None,
                            subject=title, original_text=m["original"],
                            discussion=f"移轉候選議案（第{i}次）；決議內容須由價格異動歷程人工拆解")
            # 專家會議僅為專家建議，不當作會議決議（B14）
            add_decision(ctx, agenda_id=ag, nature="專家建議" if m["type"] == "專家" else "會議決議", result="待確認",
                         result_text="歷史移轉：結果待人工依原文確認")
        if p.get("not_recognized"):
            result["recognition"] = None
        elif p.get("category_std") or p.get("rule_ast"):
            if merge_into_recognition_id:
                rec_id = merge_into_recognition_id
            else:
                name = title if not ctx.con.execute("SELECT 1 FROM recognition WHERE name=?", (title,)).fetchone() \
                    else f"{title}（移轉第{row['row_no']}列）"
                rec_id = create_recognition(ctx, name=name, description=f"歷史移轉：{src}")
            rule = rule_override if rule_override is not None else p.get("rule_ast")
            scope = "；".join(x for x in [p.get("rule") and f"篩選原則：{p['rule']}",
                                          p.get("old_rule") and f"原篩選條件（歷史）：{p['old_rule']}",
                                          p.get("group_code") and f"分組：{p['group_code']} {p.get('group_name') or ''}",
                                          p.get("ingredient") and f"成分/含量：{p['ingredient']}",
                                          p.get("dosage") and f"劑型：{p['dosage']}"] if x)
            cat = p["category_std"] if p["category_std"] != "取消認定" else "待確認"
            rev_id = create_revision(
                ctx, rec_id, category=cat, original_category=p.get("category"), scope_text=scope, rule=rule,
                conditions_text=p.get("side_decision"), effective_from=p.get("eff_date_iso"), case_id=case_id,
                completeness="不完整")
            result["recognition_id"] = rec_id
            result["revision_id"] = rev_id
            if p.get("cancel") or "取消候選" in p.get("category_flags", []):
                result["cancel_revision_id"] = create_revision(
                    ctx, rec_id, category="取消認定", original_category=p.get("category"),
                    scope_text="取消候選（歷史移轉）：須補依據與日期", case_id=case_id, completeness="不完整")
        ctx.con.execute("UPDATE migration_row SET status='已寫入', committed_case_id=?, committed_recognition_id=?,"
                        " decided_by=?, decided_at=? WHERE id=?",
                        (case_id, result.get("recognition_id"), ctx.actor, utc_now(), row_id))
        audit(ctx, "migration_row", row_id, "寫入", after=result)
    return result


def skip_row(ctx: Ctx, row_id, reason):
    ctx.require(ROLE_CLERK, ROLE_REVIEWER)
    if not reason:
        raise BizError("略過須填理由。")
    with ctx.tx():
        row = get(ctx, "migration_row", row_id)
        if row["status"] != "待核對":
            raise BizError("只有待核對列可略過。")
        ctx.con.execute("UPDATE migration_row SET status='略過', decided_by=?, decided_at=? WHERE id=?",
                        (ctx.actor, utc_now(), row_id))
        audit(ctx, "migration_row", row_id, "略過", reason=reason)
