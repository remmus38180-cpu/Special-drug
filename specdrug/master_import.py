"""主檔匯入：讀檔、欄位映射、試匯入檢核、原子啟用、版本差異。

不依固定欄序；使用者選工作表、標題列及映射。新快照保留舊版。
"""
from __future__ import annotations

import csv
import datetime as _dt
import json
import os
from collections import defaultdict
from pathlib import Path

from .core import BizError, Ctx, ROLE_ADMIN, ROLE_CLERK, ROLE_REVIEWER, audit
from .util import dec_text, new_id, normalize_code, parse_date, sha256_file, utc_now

EXCEL_MAX_ROW = 1048576

# 映射目標欄位：key -> 中文
MASTER_FIELDS = {
    "drug_code": "藥品代碼（必要）",
    "drug_name": "品名",
    "firm": "藥商名稱",
    "firm_code": "藥商代碼",
    "group_code": "分組代碼",
    "group_name": "分組名稱",
    "atc_code": "ATC碼",
    "atc_name": "ATC名稱",
    "dose_code": "核價劑型代碼",
    "dose_name": "核價劑型名稱",
    "ingredient": "成分",
    "strength": "含量",
    "multi_mark": "單複方註記",
    "spec_qty": "規格數值",
    "spec_unit": "規格單位",
    "price": "支付價",
    "annotation_raw": "必要藥品註記（DRUG_NECESSARY）",
    "license_no": "許可證號",
}
CODE_FIELDS = {"group_code", "atc_code", "dose_code", "firm_code"}
COMPARE_FIELDS = ["drug_name", "firm", "firm_code", "group_code", "group_name", "atc_code", "dose_code",
                  "dose_name", "ingredient", "strength", "multi_mark", "spec_qty", "spec_unit", "price",
                  "annotation_raw"]

# 【設計建議】預設映射；匯入時仍由使用者確認，不視為固定欄序。
PRESETS = {
    "dwm_drug 月產製處理檔（如 dwm_drug_11509.xlsx）": {
        "mapping": {
            "drug_code": "CODE", "drug_name": "藥品名稱", "firm": "藥商名稱", "firm_code": "藥商代碼",
            "group_code": "分組代碼", "group_name": "分組名稱", "atc_code": "ATC7碼", "atc_name": "ATC名稱",
            "dose_code": "PDOSE", "dose_name": "核價劑型", "ingredient": "成分名稱", "strength": "成分含量",
            "multi_mark": "複方藥品", "spec_qty": "規格量", "spec_unit": "規格單位",
            "price": "price11510", "annotation_raw": "必要藥品", "license_no": "許可證號",
        },
        "lookups": {},
        "note": "支付價欄請依資料基準日選 PRICE11509／price11510 等。",
    },
    "DWM_DRUG SQL Server每日檔（依table layout欄名）": {
        "mapping": {
            "drug_code": "DRUG_CODE", "drug_name": "DRUG_ENAME", "firm_code": "DRUGGIST_ID",
            "group_code": "DRUG_GROUP_CODE", "atc_code": "DRUG_ATC7_CODE", "dose_code": "DRUG_APV_DOSE_CODE",
            "ingredient": "DRUG_ING_CODE", "strength": "DRUG_ING_QTY", "multi_mark": "DRUG_MULTI_MARK",
            "spec_qty": "DRUG_STD_QTY_CODE", "spec_unit": "DRUG_STD_CODE",
            "annotation_raw": "DRUG_NECESSARY", "license_no": "LICID_7",
        },
        # 代碼轉名稱／數值：需另選 DWM_DRUG_RCODE 檔
        "lookups": {
            "spec_qty": {"type": "5", "field": "std_qty"},
            "spec_unit": {"type": "2", "field": "ename"},
            "group_name": {"type": "8", "field": "cname", "from": "group_code"},
            "atc_name": {"type": "7", "field": "cname", "from": "atc_code"},
            "dose_name": {"type": "1", "field": "cname", "from": "dose_code"},
        },
        "note": "支付價另選 DWM_DRPAY 長表並指定資料基準日；規格量／單位代碼經 DWM_DRUG_RCODE 轉換。",
    },
}


# ---------------- 讀檔 ----------------

def list_sheets(path) -> list[str]:
    if str(path).lower().endswith(".csv"):
        return ["(CSV)"]
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def iter_rows(path, sheet=None, header_row=1):
    """產生 (headers, iterator of (excel_row_no, dict))。CSV一律文字讀入，保留前導零。"""
    p = str(path)
    if p.lower().endswith(".csv"):
        f = open(p, "r", encoding="utf-8-sig", newline="")
        try:
            reader = csv.reader(f)
            rows = enumerate(reader, start=1)
            headers = None
            for no, r in rows:
                if no == header_row:
                    headers = [h.strip() for h in r]
                    break
        except UnicodeDecodeError:
            f.close()
            raise BizError("CSV非UTF-8編碼，請另存為「CSV UTF-8」後再匯入。")
        if headers is None:
            f.close()
            raise BizError("找不到標題列。")

        def gen():
            try:
                for no, r in rows:
                    if not any(c.strip() for c in r):
                        continue
                    yield no, dict(zip(headers, r))
            except UnicodeDecodeError:
                raise BizError("CSV非UTF-8編碼，請另存為「CSV UTF-8」後再匯入。")
            finally:
                f.close()
        return headers, gen()
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    it = ws.iter_rows(min_row=header_row, values_only=True)
    try:
        headers = [("" if h is None else str(h).strip()) for h in next(it)]
    except StopIteration:
        wb.close()
        raise BizError("工作表沒有資料。")

    def gen():
        try:
            for no, r in enumerate(it, start=header_row + 1):
                if r is None or all(c is None or (isinstance(c, str) and not c.strip()) for c in r):
                    continue
                yield no, dict(zip(headers, r))
        finally:
            wb.close()
    return headers, gen()


def headers_of(path, sheet=None, header_row=1) -> list[str]:
    h, g = iter_rows(path, sheet, header_row)
    g.close()
    return h


def _text(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (_dt.datetime, _dt.date)):
        return parse_date(v)
    t = str(v).strip()
    return t or None


# ---------------- 輔助檔 ----------------

def load_rcode(path) -> dict:
    """DWM_DRUG_RCODE：{(type, code): {cname, ename, std_qty}}"""
    _, rows = iter_rows(path)
    out = {}
    for _, r in rows:
        t = _text(r.get("DRUG_REL_CODE_TYPE"))
        c = _text(r.get("DRUG_REL_CODE"))
        if not t or not c:
            continue
        out[(t, c)] = {"cname": _text(r.get("DRUG_REL_CODE_CNAME")), "ename": _text(r.get("DRUG_REL_CODE_ENAME")),
                       "std_qty": dec_text(r.get("DRUG_STD_QTY"))}
    if not out:
        raise BizError("代碼檔沒有可用資料，請確認是 DWM_DRUG_RCODE 匯出檔（需有 DRUG_REL_CODE_TYPE 等欄位）。")
    return out


def load_price_long(path, base_date: str, code_col="DRUG_CODE", s_col="PAY_S_DATE", e_col="PAY_E_DATE",
                    price_col="PAY_PRICE"):
    """多期價格長表：依基準日取 起日<=基準日<=迄日 的價格；同代碼多筆不同價→矛盾。"""
    if not base_date:
        raise BizError("價格長表須先指定資料基準日，不任取第一筆。")
    headers, rows = iter_rows(path)
    for c in (code_col, s_col, e_col, price_col):
        if c not in headers:
            rows.close()
            raise BizError(f"價格檔缺少欄位 {c}，請確認欄位映射。")
    prices = {}
    conflicts = {}
    last_no = 0
    for no, r in rows:
        last_no = no
        code, err = normalize_code(r.get(code_col))
        if err:
            continue
        s = parse_date(r.get(s_col))
        e = parse_date(r.get(e_col))
        if not s or s > base_date or (e and e < base_date):
            continue
        p = dec_text(r.get(price_col))
        if code in prices and prices[code] != p:
            conflicts.setdefault(code, {prices[code]}).add(p)
        prices[code] = p
    return prices, {k: sorted(map(str, v)) for k, v in conflicts.items()}, last_no >= EXCEL_MAX_ROW


# ---------------- 試匯入 ----------------

def stage_import(ctx: Ctx, path, mapping: dict, sheet=None, header_row=1, base_date=None,
                 lookups: dict | None = None, rcode_path=None, price_path=None, note=None) -> str:
    """試匯入：寫入暫存版本（狀態=試匯入），回傳import_id；不改目前主檔版本。"""
    ctx.require(ROLE_ADMIN, ROLE_REVIEWER, ROLE_CLERK)
    if "drug_code" not in mapping or not mapping["drug_code"]:
        raise BizError("必須映射「藥品代碼」欄位。")
    bad = set(mapping) - set(MASTER_FIELDS)
    if bad:
        raise BizError(f"未知的映射目標：{', '.join(sorted(bad))}")
    if base_date and not parse_date(base_date):
        raise BizError("資料基準日格式錯誤。")
    base_date = parse_date(base_date) if base_date else None
    lookups = lookups or {}
    if lookups and not rcode_path:
        raise BizError("此映射需代碼轉換，請另選 DWM_DRUG_RCODE 檔。")
    rcode = load_rcode(rcode_path) if rcode_path else {}

    headers, rows = iter_rows(path, sheet, header_row)
    missing_src = [f"{MASTER_FIELDS[t]}←{s}" for t, s in mapping.items() if s and s not in headers]
    if missing_src:
        rows.close()
        raise BizError("來源檔找不到映射欄位：" + "、".join(missing_src))

    price_map, price_conflicts, price_trunc = ({}, {}, False)
    if price_path:
        price_map, price_conflicts, price_trunc = load_price_long(price_path, base_date)

    sha = sha256_file(path)
    import_id = new_id()
    records: dict[str, dict] = {}
    rows_of: dict[str, list[int]] = defaultdict(list)
    code_errors, warnings = [], []
    identical_dups, conflicts = [], []
    unknown_lookup = defaultdict(set)
    last_no = 0
    total = 0
    for no, r in rows:
        last_no = no
        total += 1
        raw_code = r.get(mapping["drug_code"])
        code, err = normalize_code(raw_code)
        if err:
            code_errors.append({"row": no, "error": err})
            continue
        rec = {"drug_code_raw": None if raw_code is None else str(raw_code)}
        for target, src in mapping.items():
            if target == "drug_code" or not src:
                continue
            v = r.get(src)
            if target in CODE_FIELDS and isinstance(v, (int, float)) and not isinstance(v, bool):
                warnings.append(f"第{no}列 {MASTER_FIELDS[target]} 為數值型（{v}），可能遺失前導零")
            if target == "price":
                rec[target] = dec_text(v)
            elif target == "spec_qty" and target not in lookups:
                rec[target] = dec_text(v) if v is not None and str(v).strip() else None
            else:
                rec[target] = _text(v)
        for target, lk in lookups.items():
            src_target = lk.get("from", target)
            code_val = rec.get(src_target)
            if code_val is None:
                if lk.get("from") is None:
                    rec[target] = None
                continue
            hit = rcode.get((lk["type"], code_val))
            if hit is None:
                unknown_lookup[target].add(code_val)
                if lk.get("from") is None:
                    rec[target] = None
            else:
                rec[target] = hit.get(lk["field"])
        if price_path:
            rec["price"] = price_map.get(code)
        rec["source_row"] = no
        rows_of[code].append(no)
        if code in records:
            prev = records[code]
            diffs = {k: [prev.get(k), rec.get(k)] for k in COMPARE_FIELDS + ["license_no"]
                     if prev.get(k) != rec.get(k)}
            if diffs:
                conflicts.append({"code": code, "rows": list(rows_of[code]), "diffs": diffs})
            else:
                identical_dups.append({"code": code, "rows": list(rows_of[code])})
            continue
        records[code] = rec

    for code, vals in price_conflicts.items():
        if code in records:
            conflicts.append({"code": code, "rows": rows_of[code], "diffs": {"price": vals},
                              "note": "價格長表於基準日有多筆不同價格"})
    if last_no >= EXCEL_MAX_ROW or price_trunc:
        warnings.insert(0, "來源資料達Excel列數上限（1,048,576列），可能已被截斷，請改用完整CSV匯出。")

    unmapped = [MASTER_FIELDS[t] for t in MASTER_FIELDS if t not in mapping and t not in lookups
                and not (t == "price" and price_path)]
    check = {
        "total_rows": total, "codes": len(records),
        "code_errors": code_errors[:500], "code_error_count": len(code_errors),
        "identical_duplicates": identical_dups[:500], "identical_duplicate_count": len(identical_dups),
        "conflicts": conflicts[:500], "conflict_count": len(conflicts),
        "unmapped_fields": unmapped,
        "unknown_lookup": {k: sorted(v)[:50] for k, v in unknown_lookup.items()},
        "warnings": warnings[:200], "warning_count": len(warnings),
        "base_date": base_date,
        "price_source": os.path.basename(str(price_path)) if price_path else None,
    }
    conflict_codes = {c["code"] for c in conflicts}
    with ctx.tx():
        ctx.con.execute(
            "INSERT INTO master_import(id, kind, source_name, sha256, base_date, sheet, header_row, mapping_json,"
            " imported_by, imported_at, row_count, check_json, status, note) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (import_id, "drug", os.path.basename(str(path)), sha, base_date, sheet, header_row,
             json.dumps({"mapping": mapping, "lookups": lookups,
                         "rcode": os.path.basename(str(rcode_path)) if rcode_path else None,
                         "price": os.path.basename(str(price_path)) if price_path else None},
                        ensure_ascii=False),
             ctx.actor, utc_now(), len(records), json.dumps(check, ensure_ascii=False), "試匯入", note))
        cols = ["import_id", "drug_code", "drug_code_raw", *[t for t in MASTER_FIELDS if t != "drug_code"],
                "source_row", "extra_json"]
        sql = f"INSERT INTO master_drug({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})"
        batch = []
        for code, rec in records.items():
            extra = {"conflict": True} if code in conflict_codes else None
            batch.append([import_id, code, rec.get("drug_code_raw"),
                          *[rec.get(t) for t in MASTER_FIELDS if t != "drug_code"],
                          rec["source_row"], json.dumps(extra) if extra else None])
        ctx.con.executemany(sql, batch)
        audit(ctx, "master_import", import_id, "試匯入",
              after={"source": os.path.basename(str(path)), "sha256": sha, "codes": len(records),
                     "conflicts": len(conflicts)})
    return import_id


def check_of(con, import_id) -> dict:
    row = con.execute("SELECT check_json FROM master_import WHERE id=?", (import_id,)).fetchone()
    if not row:
        raise BizError("找不到主檔匯入版本。")
    return json.loads(row[0])


def activate(ctx: Ctx, import_id: str, confirm_merge_duplicates=False) -> None:
    """原子性啟用；有矛盾列阻止。失敗不改目前主檔版本。"""
    ctx.require(ROLE_ADMIN, ROLE_REVIEWER)
    imp = ctx.con.execute("SELECT * FROM master_import WHERE id=?", (import_id,)).fetchone()
    if not imp:
        raise BizError("找不到主檔匯入版本。")
    if imp["status"] != "試匯入":
        raise BizError(f"此版本狀態為「{imp['status']}」，不能再啟用。")
    chk = json.loads(imp["check_json"])
    if chk["conflict_count"]:
        rows = "；".join(f"{c['code']}（列 {','.join(map(str, c['rows']))}：{','.join(c['diffs'])}）"
                        for c in chk["conflicts"][:10])
        raise BizError(f"同代碼內容矛盾 {chk['conflict_count']} 筆，阻止啟用：{rows}")
    if chk["identical_duplicate_count"] and not confirm_merge_duplicates:
        raise BizError(f"有 {chk['identical_duplicate_count']} 個代碼完全相同重複列，請確認合併後再啟用。")
    if chk["code_error_count"]:
        raise BizError(f"有 {chk['code_error_count']} 列藥品代碼錯誤（如數值型、空白），請修正來源後重新匯入。")
    with ctx.tx():
        prev = active_import_id(ctx.con)
        ctx.con.execute("UPDATE master_import SET status='已啟用', activated_by=?, activated_at=? WHERE id=?",
                        (ctx.actor, utc_now(), import_id))
        ctx.con.execute("INSERT INTO settings(key, value) VALUES ('active_master_import', ?)"
                        " ON CONFLICT(key) DO UPDATE SET value=excluded.value", (import_id,))
        audit(ctx, "master_import", import_id, "啟用", {"active": prev}, {"active": import_id})


def discard(ctx: Ctx, import_id: str) -> None:
    ctx.require(ROLE_ADMIN, ROLE_REVIEWER, ROLE_CLERK)
    with ctx.tx():
        imp = ctx.con.execute("SELECT status FROM master_import WHERE id=?", (import_id,)).fetchone()
        if not imp or imp[0] != "試匯入":
            raise BizError("只有試匯入版本可放棄。")
        ctx.con.execute("UPDATE master_import SET status='失敗', note=COALESCE(note,'')||' 使用者放棄' WHERE id=?",
                        (import_id,))
        ctx.con.execute("DELETE FROM master_drug WHERE import_id=?", (import_id,))
        audit(ctx, "master_import", import_id, "放棄試匯入")


def active_import_id(con) -> str | None:
    row = con.execute("SELECT value FROM settings WHERE key='active_master_import'").fetchone()
    return row[0] if row and row[0] else None


def mapped_fields(con, import_id) -> set:
    """此版本實際可用的規則欄位；未映射者規則停用。"""
    m = json.loads(con.execute("SELECT mapping_json FROM master_import WHERE id=?", (import_id,)).fetchone()[0])
    s = set(m["mapping"]) | set(m.get("lookups", {}))
    if m.get("price"):
        s.add("price")
    return s


def diff_imports(con, old_id: str, new_id_: str, fields=None) -> dict:
    fields = fields or COMPARE_FIELDS
    old = {r["drug_code"]: dict(r) for r in con.execute("SELECT * FROM master_drug WHERE import_id=?", (old_id,))}
    new = {r["drug_code"]: dict(r) for r in con.execute("SELECT * FROM master_drug WHERE import_id=?", (new_id_,))}
    added = sorted(set(new) - set(old))
    removed = sorted(set(old) - set(new))
    changed = []
    for code in sorted(set(old) & set(new)):
        d = {f: [old[code].get(f), new[code].get(f)] for f in fields if old[code].get(f) != new[code].get(f)}
        if d:
            changed.append({"drug_code": code, "diffs": d})
    return {"added": added, "removed": removed, "changed": changed}


def options(con, import_id, field) -> list[tuple[str, str]]:
    """由主檔產生選項（代碼, 說明）。"""
    pairs = {
        "group_code": ("group_code", "group_name"), "atc_code": ("atc_code", "atc_name"),
        "dose_code": ("dose_code", "dose_name"), "multi_mark": ("multi_mark", None),
        "drug_code": ("drug_code", "drug_name"), "firm_code": ("firm_code", "firm"),
        "spec_unit": ("spec_unit", None),
    }
    code_col, name_col = pairs[field]
    name_expr = f"MAX({name_col})" if name_col else "NULL"
    sql = (f"SELECT {code_col} AS c, {name_expr} AS n, COUNT(*) AS k FROM master_drug WHERE import_id=?"
           f" AND {code_col} IS NOT NULL GROUP BY {code_col} ORDER BY {code_col}")
    return [(r["c"], f"{r['n'] or ''}（{r['k']}筆）".strip()) for r in con.execute(sql, (import_id,))]
